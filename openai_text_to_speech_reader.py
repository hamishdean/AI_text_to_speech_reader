# -*- coding: utf-8 -*-


import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import threading
import queue
import concurrent.futures
import os
import tempfile
import docx
import PyPDF2
import pygame
from openai import OpenAI, AuthenticationError, APIConnectionError

# Maximum characters per TTS API request
BATCH_CHAR_LIMIT = 4000


def split_text_into_batches(text, limit=BATCH_CHAR_LIMIT):
    """Split text into batches that respect sentence boundaries where possible."""
    batches = []
    while text:
        if len(text) <= limit:
            batches.append(text)
            break
        # Try to split at the last sentence-ending punctuation within the limit
        chunk = text[:limit]
        split_pos = -1
        for sep in ['. ', '.\n', '! ', '!\n', '? ', '?\n']:
            pos = chunk.rfind(sep)
            if pos > split_pos:
                split_pos = pos + 1  # include the punctuation
        if split_pos <= 0:
            # Fall back to last space
            split_pos = chunk.rfind(' ')
        if split_pos <= 0:
            # No good break point, hard split
            split_pos = limit
        part = text[:split_pos].strip()
        if part:
            batches.append(part)
        text = text[split_pos:].strip()
    return batches


class TTSApp:
    def __init__(self, root):
        self.root = root
        self.root.title("OpenAI Text-to-Speech Reader")
        self.root.geometry("700x700")
        self.root.minsize(500, 500)

        # Initialize pygame mixer for audio playback
        pygame.mixer.init()

        # Variables
        self.api_key_var = tk.StringVar()
        self.voice_var = tk.StringVar(value="alloy")
        self.model_var = tk.StringVar(value="tts-1")
        self.speed_var = tk.StringVar(value="1.0x")
        self.concurrency_var = tk.StringVar(value="3")
        self.voices = ["alloy", "echo", "fable", "onyx", "nova", "shimmer"]
        self.models = ["tts-1", "tts-1-hd"]
        self.speeds = ["1.0x", "1.25x", "1.5x", "1.75x", "2.0x", "2.5x", "3.0x", "4.0x"]
        self.concurrency_options = ["1", "2", "3", "4", "5"]
        self.stop_requested = False
        self.is_processing = False
        self.batch_temp_files = []

        self.create_widgets()

    def create_widgets(self):
        # Main Frame
        main_frame = ttk.Frame(self.root, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)

        # --- Top Section: Settings ---
        settings_frame = ttk.LabelFrame(main_frame, text="Settings", padding="10")
        settings_frame.pack(fill=tk.X, pady=(0, 10))

        ttk.Label(settings_frame, text="OpenAI API Key:").grid(row=0, column=0, sticky=tk.W, pady=5)
        api_entry = ttk.Entry(settings_frame, textvariable=self.api_key_var, show="*", width=50)
        api_entry.grid(row=0, column=1, columnspan=3, sticky=tk.EW, padx=5, pady=5)

        ttk.Label(settings_frame, text="Voice:").grid(row=1, column=0, sticky=tk.W, pady=5)
        voice_dropdown = ttk.Combobox(settings_frame, textvariable=self.voice_var, values=self.voices, state="readonly", width=15)
        voice_dropdown.grid(row=1, column=1, sticky=tk.W, padx=5, pady=5)

        ttk.Label(settings_frame, text="Model:").grid(row=1, column=2, sticky=tk.W, padx=(15, 0), pady=5)
        model_dropdown = ttk.Combobox(settings_frame, textvariable=self.model_var, values=self.models, state="readonly", width=10)
        model_dropdown.grid(row=1, column=3, sticky=tk.W, padx=5, pady=5)

        ttk.Label(settings_frame, text="Speed:").grid(row=2, column=0, sticky=tk.W, pady=5)
        speed_dropdown = ttk.Combobox(settings_frame, textvariable=self.speed_var, values=self.speeds, state="readonly", width=10)
        speed_dropdown.grid(row=2, column=1, sticky=tk.W, padx=5, pady=5)

        ttk.Label(settings_frame, text="Parallel:").grid(row=2, column=2, sticky=tk.W, padx=(15, 0), pady=5)
        concurrency_dropdown = ttk.Combobox(settings_frame, textvariable=self.concurrency_var, values=self.concurrency_options, state="readonly", width=5)
        concurrency_dropdown.grid(row=2, column=3, sticky=tk.W, padx=5, pady=5)

        settings_frame.columnconfigure(1, weight=1)

        # --- Middle Section: Text Area ---
        text_frame = ttk.LabelFrame(main_frame, text="Text Content", padding="10")
        text_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 10))

        # Toolbar for text area
        toolbar = ttk.Frame(text_frame)
        toolbar.pack(fill=tk.X, pady=(0, 5))

        ttk.Button(toolbar, text="Load PDF", command=self.load_pdf).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(toolbar, text="Load DOCX", command=self.load_docx).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(toolbar, text="Clear Text", command=self.clear_text).pack(side=tk.LEFT)

        # Text Widget with Scrollbar
        self.text_area = tk.Text(text_frame, wrap=tk.WORD, font=("Segoe UI", 10))
        scrollbar = ttk.Scrollbar(text_frame, command=self.text_area.yview)
        self.text_area.configure(yscrollcommand=scrollbar.set)

        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.text_area.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # --- Batch Progress Section ---
        batch_frame = ttk.LabelFrame(main_frame, text="Batch Progress", padding="10")
        batch_frame.pack(fill=tk.X, pady=(0, 10))

        self.batch_progress_var = tk.StringVar(value="No batches to process.")
        self.batch_progress_label = ttk.Label(batch_frame, textvariable=self.batch_progress_var, wraplength=650)
        self.batch_progress_label.pack(fill=tk.X)

        # Progress bar
        self.progress_bar = ttk.Progressbar(batch_frame, mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=(5, 0))

        # Batch log text box
        self.batch_log = tk.Text(batch_frame, height=4, wrap=tk.WORD, font=("Segoe UI", 9), state=tk.DISABLED)
        batch_scroll = ttk.Scrollbar(batch_frame, command=self.batch_log.yview)
        self.batch_log.configure(yscrollcommand=batch_scroll.set)
        batch_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.batch_log.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, pady=(5, 0))

        # --- Bottom Section: Controls ---
        control_frame = ttk.Frame(main_frame)
        control_frame.pack(fill=tk.X, pady=(5, 0))

        self.play_btn = ttk.Button(control_frame, text="▶ Read Aloud", command=self.start_reading, style="Accent.TButton")
        self.play_btn.pack(side=tk.LEFT, padx=(0, 5))

        self.stop_btn = ttk.Button(control_frame, text="⏹ Stop", command=self.stop_audio, state=tk.DISABLED)
        self.stop_btn.pack(side=tk.LEFT)

        self.status_var = tk.StringVar(value="Ready.")
        status_label = ttk.Label(control_frame, textvariable=self.status_var, foreground="gray")
        status_label.pack(side=tk.RIGHT)

    def log_batch(self, message):
        """Append a message to the batch log box."""
        self.batch_log.config(state=tk.NORMAL)
        self.batch_log.insert(tk.END, message + "\n")
        self.batch_log.see(tk.END)
        self.batch_log.config(state=tk.DISABLED)

    def clear_batch_log(self):
        """Clear the batch log box."""
        self.batch_log.config(state=tk.NORMAL)
        self.batch_log.delete(1.0, tk.END)
        self.batch_log.config(state=tk.DISABLED)

    def load_pdf(self):
        file_path = filedialog.askopenfilename(filetypes=[("PDF Files", "*.pdf")])
        if not file_path:
            return

        self.status_var.set("Extracting text from PDF...")
        self.root.update()

        try:
            text = ""
            with open(file_path, "rb") as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"

            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, text)
            self.status_var.set(f"Loaded {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read PDF:\n{str(e)}")
            self.status_var.set("Error loading file.")

    def load_docx(self):
        file_path = filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
        if not file_path:
            return

        self.status_var.set("Extracting text from DOCX...")
        self.root.update()

        try:
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])

            self.text_area.delete(1.0, tk.END)
            self.text_area.insert(tk.END, text)
            self.status_var.set(f"Loaded {os.path.basename(file_path)}")
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read DOCX:\n{str(e)}")
            self.status_var.set("Error loading file.")

    def clear_text(self):
        self.text_area.delete(1.0, tk.END)
        self.status_var.set("Ready.")

    def stop_audio(self):
        self.stop_requested = True
        if pygame.mixer.music.get_busy():
            pygame.mixer.music.stop()
        self.play_btn.config(state=tk.NORMAL)
        self.stop_btn.config(state=tk.DISABLED)
        self.status_var.set("Stopped.")
        self.log_batch("Stopped by user.")

    def start_reading(self):
        # Guard against starting a second batch while one is running
        if self.is_processing:
            return

        api_key = self.api_key_var.get().strip()
        text = self.text_area.get(1.0, tk.END).strip()

        if not api_key:
            messagebox.showwarning("Missing API Key", "Please enter your OpenAI API key.")
            return
        if not text:
            messagebox.showwarning("Empty Text", "There is no text to read.")
            return

        # Parse speed value
        speed_str = self.speed_var.get().replace("x", "")
        try:
            speed = float(speed_str)
        except ValueError:
            speed = 1.0

        # Split text into batches
        batches = split_text_into_batches(text)
        total = len(batches)

        if total == 0:
            messagebox.showwarning("Empty Text", "There is no text to read.")
            return

        # Update UI
        self.clear_batch_log()
        self.play_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        self.stop_requested = False
        self.is_processing = True
        self.progress_bar['value'] = 0
        self.progress_bar['maximum'] = total

        if total == 1:
            self.batch_progress_var.set("1 batch to process (text fits in a single request).")
        else:
            self.batch_progress_var.set(f"{total} batches to process.")
        self.log_batch(f"Text split into {total} batch(es) ({len(text)} characters total).")

        model = self.model_var.get()
        voice = self.voice_var.get()

        # Parse concurrency
        try:
            max_workers = int(self.concurrency_var.get())
        except ValueError:
            max_workers = 3

        # Queue for producer (generator) -> consumer (player) communication
        # Items are (batch_num, temp_file) tuples, or None as a sentinel for "done"
        self.audio_queue = queue.Queue()
        self.generator_error = None

        # Launch producer coordinator and consumer threads
        threading.Thread(
            target=self.generate_batches_concurrent,
            args=(api_key, batches, voice, model, speed, total, max_workers),
            daemon=True
        ).start()
        threading.Thread(
            target=self.play_batches,
            args=(total,),
            daemon=True
        ).start()

    def generate_single_batch(self, client, batch_index, batch_text, voice, model, speed, total):
        """Worker: generate audio for a single batch. Returns (batch_num, temp_file)."""
        batch_num = batch_index + 1
        self.root.after(0, lambda n=batch_num, chars=len(batch_text): self.log_batch(
            f"Batch {n}/{total}: Generating audio ({chars} chars)..."
        ))

        temp_dir = tempfile.gettempdir()
        temp_file = os.path.join(temp_dir, f"tts_batch_{os.getpid()}_{batch_index}.mp3")

        response = client.audio.speech.create(
            model=model,
            voice=voice,
            input=batch_text,
            speed=speed,
        )
        response.stream_to_file(temp_file)

        self.root.after(0, lambda n=batch_num: self.log_batch(
            f"Batch {n}/{total}: Audio generated, ready to play."
        ))

        return batch_num, temp_file

    def generate_batches_concurrent(self, api_key, batches, voice, model, speed, total, max_workers):
        """Coordinator: submit all batches to a thread pool, then feed results to
        the playback queue in order."""
        try:
            client = OpenAI(api_key=api_key)
        except Exception as e:
            self.generator_error = f"Failed to create client:\n{str(e)}"
            self.audio_queue.put(None)
            return

        generated_count = 0

        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
            # Submit all batches to the pool
            futures = {}
            for i, batch_text in enumerate(batches):
                if self.stop_requested:
                    break
                future = executor.submit(
                    self.generate_single_batch,
                    client, i, batch_text, voice, model, speed, total
                )
                futures[i] = future

            # Collect results in order so playback stays sequential
            for i in range(len(batches)):
                if self.stop_requested:
                    # Cancel any pending futures
                    for f in futures.values():
                        f.cancel()
                    self.audio_queue.put(None)
                    return

                if i not in futures:
                    break

                try:
                    batch_num, temp_file = futures[i].result()
                    self.batch_temp_files.append(temp_file)
                    generated_count += 1
                    self.root.after(0, lambda n=generated_count: self.progress_bar.configure(value=n))
                    self.audio_queue.put((batch_num, temp_file))
                except AuthenticationError:
                    self.generator_error = "Invalid OpenAI API Key."
                    for f in futures.values():
                        f.cancel()
                    self.audio_queue.put(None)
                    return
                except APIConnectionError:
                    self.generator_error = "Network error. Please check your connection."
                    for f in futures.values():
                        f.cancel()
                    self.audio_queue.put(None)
                    return
                except Exception as e:
                    self.generator_error = f"Error on batch {i + 1}:\n{str(e)}"
                    for f in futures.values():
                        f.cancel()
                    self.audio_queue.put(None)
                    return

        # Signal that all batches have been generated
        self.audio_queue.put(None)

    def play_batches(self, total):
        """Consumer: play audio files from the queue as they become available."""
        while True:
            if self.stop_requested:
                self.root.after(0, lambda: self.log_batch("Playback stopped."))
                self.root.after(0, self.reset_ui)
                return

            # Wait for next item (blocks until producer puts something)
            try:
                item = self.audio_queue.get(timeout=0.5)
            except queue.Empty:
                continue

            # None sentinel means producer is done (either finished or errored)
            if item is None:
                if self.generator_error:
                    err_msg = self.generator_error
                    self.root.after(0, lambda msg=err_msg: messagebox.showerror("Error", msg))
                    self.root.after(0, self.reset_ui)
                else:
                    # All batches generated and played
                    self.root.after(0, lambda: self.status_var.set("Ready."))
                    self.root.after(0, lambda: self.batch_progress_var.set(f"All {total} batch(es) completed."))
                    self.root.after(0, lambda: self.log_batch("All batches finished."))
                    self.root.after(0, self.reset_ui)
                return

            batch_num, temp_file = item

            if self.stop_requested:
                self.root.after(0, self.reset_ui)
                return

            try:
                pygame.mixer.music.load(temp_file)
                pygame.mixer.music.play()
                self.root.after(0, lambda n=batch_num: self.status_var.set(f"Playing batch {n}/{total}..."))
                self.root.after(0, lambda n=batch_num: self.batch_progress_var.set(
                    f"Playing batch {n}/{total}..."
                ))

                # Wait for playback to complete
                while pygame.mixer.music.get_busy():
                    if self.stop_requested:
                        pygame.mixer.music.stop()
                        self.root.after(0, self.reset_ui)
                        return
                    pygame.time.Clock().tick(10)

            except Exception as e:
                err_msg = f"Failed to play batch {batch_num}:\n{str(e)}"
                self.root.after(0, lambda msg=err_msg: messagebox.showerror("Playback Error", msg))
                self.root.after(0, self.reset_ui)
                return

            self.root.after(0, lambda n=batch_num: self.log_batch(f"Batch {n}/{total}: Playback complete."))

    def cleanup_temp_files(self):
        """Unload pygame music and remove temporary audio files."""
        try:
            pygame.mixer.music.unload()
        except Exception:
            pass
        for f in self.batch_temp_files:
            try:
                if os.path.exists(f):
                    os.remove(f)
            except OSError:
                pass
        self.batch_temp_files = []

    def reset_ui(self):
        self.play_btn.config(state=tk.NORMAL)
        self.stop_btn.config(state=tk.DISABLED)
        self.is_processing = False
        self.cleanup_temp_files()

if __name__ == "__main__":
    root = tk.Tk()

    # Optional: Apply a slightly better default style if available
    style = ttk.Style()
    if "clam" in style.theme_names():
        style.theme_use("clam")

    app = TTSApp(root)
    root.mainloop()
